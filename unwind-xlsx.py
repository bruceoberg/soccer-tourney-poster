import openpyxl
import pyaml

wb = openpyxl.load_workbook(filename = '2022-world-cup.xlsx')
mpStrTableLMpKV = {}
for ws in wb.worksheets:
	for strTable, strRange in ws.tables.items():
		print(f'table:{strTable} ... {strRange}')
		lStrKey = None
		lMpKV = []
		for row in ws[strRange]:
			if not lStrKey:
				lStrKey = [cell.value.lower() for cell in row]
			else:
				lStrValue = [cell.value for cell in row]
				lMpKV.append(dict(zip(lStrKey, lStrValue)))
		mpStrTableLMpKV[strTable] = lMpKV

with open('2022-world-cup.yaml', 'w') as fd:
	pyaml.dump(mpStrTableLMpKV, fd)